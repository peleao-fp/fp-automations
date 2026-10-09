#!/usr/bin/env python3
"""
Ferramentas dos hardgoods — o que hoje se faz à mão, em comandos com simulação. Nada grava sem --apply.

  python -m hardgoods.cli pb 68D5F78C                         # mostra o PB (cabeçalho e linhas)
  python -m hardgoods.cli remove 68D5F78C 3203-24-60 4111-06-09 [--apply]
  python -m hardgoods.cli buy 68D5F78C 22BF77D9 --grower SSA --ship 2026-10-12 --so 68D5F78C=SO-454843 [--costs custos.json] [--apply]
  python -m hardgoods.cli order 68D5F78C 22BF77D9 --supplier Syndicate --out ~/Downloads
  python -m hardgoods.cli merge 68D5F78C 29301624 [--apply]     # junta o PB 29301624 dentro do 68D5F78C (fica o 1º)
  python -m hardgoods.cli packing --grower SSA --ship 2026-10-12 [--available 2026-10-16] [--apply]

PB é sempre o código interno (pbook_uq); a automação o guarda ao criar o PB.
"""
import argparse, json, math, os, re, sys
from datetime import date, timedelta

from . import flexy, fullpotos as FP, purchase as P

CODE = re.compile(r"\b(?:\d{2,5}-\d{2,3}-\d{2,4}|C\d{3,4}(?:-\d{2})?|[A-Z]{1,4}\d{2,5}[A-Z]{0,3})\b")
GROWERS = {"SSA": "1B5CCCF3", "SFP": "6E2B3D86", "WDI": "08A2D2A6", "GWS": None}   # prefixo do PO → grower_uq
STORE = {"FT LAUDERDALE": "PMP", "WEST PALM BEACH": "WPB", "NAPLES": "NPL", "MIAMI": "MIA"}
SKIP = ("DELIVERY CHARGE",)

def codes(desc): return CODE.findall((desc or "").upper())
def has_code(desc, c): return re.search(r"(^|[\s(/])" + re.escape(c.upper()) + r"($|[\s)])", (desc or "").upper().strip()) is not None

def pb_info(uq):
    h = flexy.read_header(uq)
    if not h: sys.exit(f"PB {uq} não encontrado")
    wh_name = (h.get("warehouse") or "").split(" - ")[0].strip().upper()
    m = re.search(r"-\s*(\d+)\s*$", (h.get("customer") or "").strip())
    return {"uq": uq, "no": h["pbook_no"], "status": h.get("status"), "customer": (h.get("customer") or "").strip(),
            "fp": f"FP#{m.group(1)}" if m else "", "store": STORE.get(wh_name, wh_name), "wh": P.WH_BY_NAME.get(wh_name),
            "date": h.get("shipping_date")}

def cmd_pb(a):
    for uq in a.pb:
        i = pb_info(uq); L = flexy.read_lines(uq)
        print(f"PB {i['no']} ({uq}) · {i['store']} · {i['customer']} · {i['status']} · {i['date']} · {len(L)} linhas, {sum(l['qty_order'] for l in L)} cx")
        for l in sorted(L, key=lambda l: l["description"]):
            print(f"   {l['description'].strip()[:48]:<48} {l['qty_order']:>3} cx × {l['packs_x_case']} × {l['up_x_pack']}"
                  f"  venda {l['so_price']:>8}  já em PO: {l['qty_porder']}")

def cmd_remove(a):
    L = flexy.read_lines(a.pb[0])
    for c in a.codes:
        m = [l for l in L if has_code(l["description"], c)]
        if len(m) != 1: print(f"  {c}: {len(m)} linhas — não mexo"); continue
        if m[0]["qty_porder"]: print(f"  {c}: já tem PO ({m[0]['qty_porder']} cx) — não mexo"); continue
        print(f"  {'remove' if a.apply else '(simulação) removeria'} {m[0]['description'].strip()[:50]} ({m[0]['qty_order']} cx)")
        if a.apply: flexy.delete_line(m[0]["unico"])
    if a.apply:
        L = flexy.read_lines(a.pb[0]); print(f"  PB agora: {len(L)} linhas, {sum(l['qty_order'] for l in L)} cx")

def last_costs(prefix, days=75):
    """Custo por unidade da última compra de cada código desse fornecedor."""
    hist = {}
    for i in range(days):
        d = (date.today() - timedelta(days=i)).isoformat()
        try: rows = P.pos_by_date(d)
        except Exception: continue
        for p in rows:
            if not (p.get("grower_po") or "").startswith(prefix): continue
            for c in codes(p.get("description")):
                if c not in hist or d > hist[c][0]: hist[c] = (d, p["unico"])
    out = {}
    for c, (d, unico) in hist.items():          # preço do PO é por pack → divide pelas unidades do pack
        po = P.po_detail(unico)
        if po: out[c] = (d, float(po.get("po_price") or 0) / max(1, int(po.get("up_x_pack") or 1)))
    return out

def cost_for(l, hist, manual):
    if l["unico"] in manual: return float(manual[l["unico"]]), "informado"
    upc = int(l["packs_x_case"]) * int(l["up_x_pack"]); sale = float(l["so_price"]) / max(1, int(l["up_x_pack"]))
    h = next((hist[c] for c in codes(l["description"]) if c in hist), None)
    if h: return h[1], f"última compra {h[0]}"
    if sale < 0.05: return None, "SEM PREÇO"
    if sale >= 10 and upc >= 12: return math.floor(sale / upc * 0.62 * 100) / 100, f"venda era da caixa ({sale:.2f}÷{upc})"
    return math.floor(sale * 0.62 * 100) / 100, "venda × 0,62"

def cmd_buy(a):
    grower = GROWERS.get(a.grower.upper())
    if not grower: sys.exit(f"Fornecedor {a.grower} sem grower_uq configurado")
    manual = json.load(open(a.costs)) if a.costs else {}
    so = dict(x.split("=", 1) for x in (a.so or []))
    print("Buscando custo das últimas compras…"); hist = last_costs(a.grower.upper())
    plan, total = [], 0.0
    for uq in a.pb:
        i = pb_info(uq)
        if not i["wh"]: sys.exit(f"PB {i['no']}: warehouse desconhecido")
        for l in flexy.read_lines(uq):
            if any(s in l["description"].upper() for s in SKIP): continue
            falta = int(l["qty_order"]) - int(l["qty_porder"] or 0)
            if falta <= 0: continue
            cost, how = cost_for(l, hist, manual)
            upc = int(l["packs_x_case"]) * int(l["up_x_pack"]); ext = round((cost or 0) * upc * falta, 2); total += ext
            plan.append((i, l, falta, cost, how))
            print(f"{'!!' if cost is None else '  '} {i['store']:<4} {i['no']} {l['description'].strip()[:40]:<40} {falta:>2} cx × {upc:>3} | "
                  f"custo/un {('%.4f' % cost) if cost is not None else '   —  '} | {how}")
    print(f"\n{len(plan)} linhas a comprar · total estimado ${total:,.2f}")
    if not a.apply: return print("(simulação — nada gravado)")
    if any(p[3] is None for p in plan): sys.exit("NÃO GRAVADO: há linhas sem custo (use --costs)")
    for i, l, falta, cost, how in plan:
        note = f"{so.get(i['uq'], '')} - " + ("custo da ultima compra" if how.startswith("última") else "custo estimado - confirmar na invoice")
        try:
            po = P.create_from_pb_line(l, i["uq"], grower, i["wh"], a.ship, cost, qty=falta, details=note.strip(" -"))
            print(f"  OK   {po['description'].strip()[:40]:<40} {po['qty_porder']} cx = ${po['ext_price']}")
        except Exception as e:
            print(f"  ERRO {l['description'].strip()[:40]:<40} {e}")

def cmd_order(a):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    os.makedirs(os.path.expanduser(a.out), exist_ok=True)
    for uq in a.pb:
        i = pb_info(uq); wb = Workbook(); ws = wb.active; ws.title = "Order"
        ws["A1"] = f"FULL POT · {a.supplier.upper()} ORDER"; ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"STORE: {i['store']} · {i['customer']}"; ws["A2"].font = Font(bold=True, color="C00000")
        ws["A3"] = "Delivery: 1516 SW 13 CT, Pompano Beach, FL 33069"
        ws.append([]); ws.append(["Box Mark", "PB", "Item", "Description", "Boxes", "Units/Box", "Total Units"])
        for l in sorted(flexy.read_lines(uq), key=lambda l: l["description"]):
            if any(s in l["description"].upper() for s in SKIP): continue
            c = codes(l["description"]); upc = int(l["packs_x_case"]) * int(l["up_x_pack"])
            ws.append([i["fp"], i["no"], c[0] if c else "", l["description"].strip(), l["qty_order"], upc, l["qty_order"] * upc])
        path = os.path.join(os.path.expanduser(a.out), f"{a.supplier} - {i['store']} - PB {i['no']}.xlsx"); wb.save(path); print("  ", path)


def cmd_merge(a):
    """Junta os PBs de origem dentro do PB de destino. Item igual (produto, formato e preço) soma; o resto entra como linha nova."""
    target, sources = a.pb[0], a.pb[1:]
    T = flexy.read_lines(target); S = {u: flexy.read_lines(u) for u in sources}
    for u, L in S.items():
        if any(int(l["qty_porder"] or 0) for l in L): sys.exit(f"PB {u} já tem linha com PO — não junto")
    os.makedirs("out", exist_ok=True)
    with open(f"out/backup_merge_{target}.json", "w") as f: json.dump({"target": T, "sources": S}, f)
    same = lambda x, l: (x["product_uq"] == l["product_uq"] and x["packs_x_case"] == l["packs_x_case"]
                         and x["up_x_pack"] == l["up_x_pack"] and abs(float(x["so_price"]) - float(l["so_price"])) < 0.001)
    plan = []   # (linha a inserir, qty, linha antiga do destino a apagar ou None)
    for u, L in S.items():
        for l in L:
            t = [x for x in T if same(x, l) and not int(x["qty_porder"] or 0)]
            plan.append((l, int(l["qty_order"]) + (int(t[0]["qty_order"]) if t else 0), t[0] if t else None))
    for l, q, old in plan:
        print(f"  {'SOMA ' if old else 'ENTRA'} {l['description'].strip()[:44]:<44} {q} cx × {l['packs_x_case']} × {l['up_x_pack']} @ {l['so_price']}")
    exp = sum(int(l["qty_order"]) for l in T) + sum(int(l["qty_order"]) for L in S.values() for l in L)
    print(f"  destino ficará com {exp} cx")
    if not a.apply: return print("(simulação — nada gravado)")
    for l, q, old in plan:   # insere tudo primeiro (pula o que já entrou, se for uma retomada)
        cur = flexy.read_lines(target)
        if any(same(x, l) and int(x["qty_order"]) == q and (not old or x["unico"] != old["unico"]) for x in cur): continue
        flexy.insert_line(target, l["product_uq"], l["case_uq"], q, int(l["packs_x_case"]), int(l["up_x_pack"]), float(l["so_price"]))
    for l, q, old in plan:
        if old: flexy.delete_line(old["unico"])
    got = sum(int(l["qty_order"]) for l in flexy.read_lines(target))
    if got != exp: sys.exit(f"CONFERÊNCIA FALHOU: destino tem {got} cx, esperado {exp} — origens NÃO apagadas")
    for u, L in S.items():
        for l in L: flexy.delete_line(l["unico"])
        flexy.delete_prebook(u)
    print(f"  OK: destino com {got} cx; origens esvaziadas")

def cmd_packing(a):
    """Fecha o packing antes da mercadoria: um por loja, invoice XINVOICE, available date, Add from PO. Fica ABERTO."""
    grower = GROWERS.get(a.grower.upper())
    today = date.today(); avail = date.fromisoformat(a.available) if a.available else today + timedelta(days=7)
    awb = "989" + avail.strftime("%Y%m%d")
    pos = [p for p in P.pos_by_date(a.ship) if (p.get("grower_po") or "").split("-")[0] == a.grower.upper() and p.get("active", True)]
    stores = {}
    for p in pos:
        falta = int(p.get("qty_porder") or 0) - int(p.get("qty_ship") or 0)   # o PO registra o que já foi para packing
        if falta <= 0: print(f"  já no Inventory Entry: {p['grower_po']} {p['description'].strip()[:40]}"); continue
        p["_falta"] = falta
        wh = P.WH_BY_NAME.get((p.get("warehouse") or "").strip().upper())
        if not wh: sys.exit(f"Warehouse desconhecido no PO {p['grower_po']}")
        stores.setdefault(wh, []).append(p)
    existing = {}   # packing já criado hoje neste AWB para esse fornecedor/warehouse (retomada)
    for pk in FP.req("GET", f"/api/inventory-entry/packing-x-awb?awb={awb}&date={today.isoformat()}") or []:
        if pk.get("grower_uq") == grower:
            h = FP.packing(pk["pack_uq"]); existing[(h or {}).get("wphysical_uq")] = pk["pack_uq"]
    print(f"{a.grower} · POs de {a.ship} · AWB {awb} · available {avail} · {'GRAVANDO' if a.apply else 'SIMULAÇÃO'}")
    for wh, L in stores.items():
        sos = sorted({m.group(0) for p in L for m in [re.search(r"SO-\d+", P.po_detail(p["unico"]).get("details") or "")] if m})
        store = next((k for k, v in P.WH_BY_NAME.items() if v == wh), wh)
        print(f"\n  {store}: {len(L)} POs, {sum(p['_falta'] for p in L)} cx {'(' + ', '.join(sos) + ')' if sos else ''}"
              f"{' — reaproveita packing ' + existing[wh] if wh in existing else ''}")
        if not a.apply: continue
        pack = existing.get(wh) or FP.create_packing(grower, wh, awb, today.strftime("%Y%m%d"), "XINVOICE", today.isoformat(), avail.isoformat(),
                                                     ("XINVOICE - AGUARDANDO INVOICE " + " ".join(sos)).strip()[:254])
        for p in L:
            try: FP.add_from_po(pack, p["unico"], p["_falta"]); print(f"    + {p['description'].strip()[:44]} ({p['_falta']} cx)")
            except Exception as e: print(f"    ERRO {p['description'].strip()[:44]}: {e}")
    if a.apply: print("\nPackings ficam ABERTOS e sem Send to WH — confira na tela do Inventory Entry.")
    else: print("\n(simulação — nada gravado)")

def cmd_invoice(a):
    """Aplica uma conferência de invoice (plano JSON): custo do PO, custo/venda da caixa, número da invoice no packing.
    Linhas 'missing' (não vieram) e de cliente (venda do vendedor) não são tocadas."""
    plan = json.load(open(a.plan)); only = set(a.only or [])
    todo = [it for it in plan if it["status"] != "missing" and it.get("box_unico") and (not only or it["box_unico"] in only)]
    print(f"{len(todo)} linhas · {'GRAVANDO' if a.apply else 'SIMULAÇÃO'}"); lost = []
    for it in todo:
        p = it["po"]; new_sale = it.get("sale_new")
        print(f"  {p['description'].strip()[:40]:<40} custo/un {it['unit_f']:.4f} (caixa {it['cost_ie']:.4f})"
              f" venda {it.get('sale_old')} → {new_sale if new_sale else 'mantém'}")
        if not a.apply: continue
        P.update(p["unico"], unit_price=it["unit_f"], confirm=True,
                 details=f"{re.sub(r' - .*$', '', p.get('details') or '').strip()} - {it['invoice']} conferida"[:250])
        if a.box_cost:   # PUT completo: muda o custo da caixa, mas ZERA o customer (bug do fullpotos)
            mark = f"FP#{p['cust'].split('-')[-1].split('/')[0].strip()}" if re.search(r"-\s*\d+\s*/", p.get("cust") or "") else None
            b, after = FP.update_box(it["box_unico"], cost=it["cost_ie"], price=new_sale, fill_box_id=mark, allow_customer_reset=True)
            if b.get("customer") != after.get("customer"):
                lost.append((b.get("lote"), p["description"].strip(), b.get("customer"), (after.get("box_id") or "").strip()))
        else:            # sem --box-cost: a caixa não é tocada (a rota "Change Prices" responde OK mas não grava)
            after = FP.box(it["box_unico"])
        print(f"     ok · BOXID '{(after.get('box_id') or '').strip()}' · customer {after.get('customer')} · custo {after.get('f_cost_x_u')} · venda {after.get('price_x_u')}")
    if lost:
        print("\nCUSTOMER A RECOLOCAR (o Edit Box zera):")
        for l in lost: print(f"  lote {l[0]} · {l[1][:40]} · customer {l[2]} · BOXID {l[3]}")
        json.dump(lost, open(os.path.splitext(a.plan)[0] + "_customers.json", "w"))
    if a.apply and a.remove_missing and not only:   # o que não veio: tira a caixa e zera o PO
        for it in plan:
            if it["status"] != "missing": continue
            p = it["po"]
            if it.get("box_unico"): FP.delete_box(it["box_unico"])
            P.update(p["unico"], qty=0, details=f"{re.sub(r' - .*$', '', p.get('details') or '').strip()} - NAO VEIO na invoice"[:250])
            print(f"  removido (não veio): {p['description'].strip()[:40]} · caixa {it.get('box_unico')} apagada · PO {p['grower_po']} zerado")
    if a.apply and a.headers and not only:
        for pk, inv in (x.split("=", 1) for x in a.headers):
            h = FP.packing(pk); det = re.sub(r"XINVOICE - AGUARDANDO INVOICE", "INVOICE " + inv, (h.get("details") or "").strip())
            h2 = FP.update_packing_header(pk, invoice_no=inv, details=det)
            print(f"  packing {pk}: invoice {(h2.get('invoice_no') or '').strip()} · {(h2.get('details') or '').strip()}")

def main():
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("pb"); s.add_argument("pb", nargs="+"); s.set_defaults(f=cmd_pb)
    s = sub.add_parser("remove"); s.add_argument("pb", nargs=1); s.add_argument("codes", nargs="+"); s.add_argument("--apply", action="store_true"); s.set_defaults(f=cmd_remove)
    s = sub.add_parser("buy"); s.add_argument("pb", nargs="+"); s.add_argument("--grower", required=True); s.add_argument("--ship", required=True)
    s.add_argument("--so", nargs="*"); s.add_argument("--costs"); s.add_argument("--apply", action="store_true"); s.set_defaults(f=cmd_buy)
    s = sub.add_parser("order"); s.add_argument("pb", nargs="+"); s.add_argument("--supplier", required=True); s.add_argument("--out", default="out"); s.set_defaults(f=cmd_order)
    s = sub.add_parser("merge"); s.add_argument("pb", nargs="+"); s.add_argument("--apply", action="store_true"); s.set_defaults(f=cmd_merge)
    s = sub.add_parser("packing"); s.add_argument("--grower", required=True); s.add_argument("--ship", required=True)
    s.add_argument("--available"); s.add_argument("--apply", action="store_true"); s.set_defaults(f=cmd_packing)
    s = sub.add_parser("invoice"); s.add_argument("--plan", required=True); s.add_argument("--only", nargs="*")
    s.add_argument("--headers", nargs="*", help="pack_uq=SI-123 ..."); s.add_argument("--box-cost", action="store_true")
    s.add_argument("--remove-missing", action="store_true"); s.add_argument("--apply", action="store_true"); s.set_defaults(f=cmd_invoice)
    a = ap.parse_args(); a.f(a)

if __name__ == "__main__":
    main()
