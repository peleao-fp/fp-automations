#!/usr/bin/env python3
"""
Pedidos das fazendas da Califórnia para o Valerii — toda quinta, olhando a sexta.

Junta, por fazenda, o que está no Purchase Control (POs com ship date de sexta) e no Inventory Entry
(caixas com data de sexta, inclusive as lançadas direto sem PO). Gera um Excel por fazenda e envia
para o Valerii pelo Resend, com cópia para o Pedro. Fazenda sem pedido não gera arquivo.

Uso:
  python orders/farm_orders.py                    # data automática (próxima sexta), envia
  python orders/farm_orders.py --date 2026-10-09  # força a data
  python orders/farm_orders.py --dry-run          # só gera os Excel em ./out
  python orders/farm_orders.py --test             # envia só para o Pedro, assunto "TESTE"

Variáveis de ambiente: RESEND_API_KEY, FULLPOTOS_COOKIE (sessão do fullpotos, para o Inventory Entry).
"""
import argparse, base64, os, sys
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from daily_orders import FROM_EMAIL, BUYER_CC, MIAMI, api_get, customer_no, ordinal

FP_BASE = "https://fullpotos.flexymax.com"
TO = ["valerii@fullpot.com"]
# prefixo do PO → (grower_uq, nome)
FARMS = {
    "ARN": ("3ACF8400", "ARANGO GREEN GROWERS"), "FGS": ("69879CED", "FLORA GREENS"),
    "KFA": ("6C936ACF", "KENDALL FARMS"),        "CAM": ("8E54BB1E", "CAMFLOR"),
    "RPT": ("A51B71D8", "RAINBOW PROTEA"),       "TFG": ("914A9740", "TWINS FLOWER FARM"),
    "OCH": ("92FED232", "OCHOA GREENS"),
}
BY_UQ = {uq: pref for pref, (uq, _) in FARMS.items()}
CASE_NAMES = {"BX": "BOX", "QB": "QUARTER", "HB": "HALF", "EB": "EIGHTH"}

def fp_get(path):
    cookie = os.environ.get("FULLPOTOS_COOKIE")
    if not cookie: raise RuntimeError("FULLPOTOS_COOKIE não configurado — sem acesso ao Inventory Entry")
    r = requests.get(FP_BASE + path, timeout=60, headers={"Cookie": cookie, "Accept": "application/json",
                     "User-Agent": "Mozilla/5.0 fp-automations"})
    if r.status_code in (401, 403): raise RuntimeError(f"fullpotos recusou o acesso (HTTP {r.status_code}) — login expirado")
    r.raise_for_status()
    return r.json()

def collect(day):
    """{prefixo: [linhas]} das 7 fazendas, juntando Purchase Control e Inventory Entry."""
    out = {p: [] for p in FARMS}
    po_units = set()
    for p in api_get(f"/api/purchase-orders?ship_date={day.isoformat()}&grower_uq=%25&product_uq=%25").get("data", []):
        pref = (p.get("grower_po") or "").split("-")[0]
        if pref not in FARMS or not p.get("active", True): continue
        d = api_get(f"/api/purchase-orders/{p['unico']}").get("data") or {}
        po_units.add(p["unico"])
        out[pref].append({
            "origem": "Purchase Control", "ref": p["grower_po"], "fp": (p.get("cporder_no") or "").strip(),
            "cliente": (p.get("customer") or "").split("/")[0].strip(), "produto": " ".join((d.get("description") or "").split()),
            "caixa": CASE_NAMES.get((d.get("case_sh") or "").strip(), (d.get("case_sh") or "").strip()),
            "qtd": int(d.get("qty_porder") or 0), "bunches": int(d.get("bunches_case") or 0), "ux": int(d.get("up_x_pack") or 0),
            "total": int(d.get("total_units") or 0), "preco": float(d.get("po_price") or 0),
            "confirmado": int(d.get("qty_confirm") or 0) >= int(d.get("qty_porder") or 0),
            "notas": " ".join((d.get("details") or "").split()).strip(". ")})
    for a in fp_get(f"/api/inventory-entry/awb-by-date?date={day.isoformat()}") or []:
        awb = (a.get("awbcode") or "").strip()
        for b in fp_get(f"/api/inventory-entry/packing-box-by-awb?awbcode={awb}") or []:
            pref = BY_UQ.get((b.get("grower_uq") or "").strip())
            if not pref or b.get("porder_uq") in po_units: continue   # já veio pelo PO
            cust = int(b.get("customer") or 0)
            out[pref].append({
                "origem": "Inventory Entry", "ref": f"Lote {b.get('lote')} · AWB {awb}", "fp": (b.get("box_id") or "").strip(),
                "cliente": f"Cliente {cust}" if cust else "ESTOQUE", "produto": " ".join((b.get("description") or "").split()),
                "caixa": CASE_NAMES.get((b.get("case_sh") or "").strip(), (b.get("case_sh") or "").strip()),
                "qtd": int(b.get("box_qty") or 0), "bunches": int(b.get("tunits_x_box") or 0), "ux": 1,
                "total": int(b.get("total_units") or 0), "preco": float(b.get("f_cost_x_u") or 0), "confirmado": True,
                "notas": " ".join((b.get("inventory_notes") or "").split())})
    return out

def build_xlsx(path, farm, day, lines):
    thin = Side(style="thin", color="999999"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook(); ws = wb.active; ws.title = "Pedidos"
    ws["A1"] = f"FULL POT  ·  {farm}"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Pedidos para {day:%A, %m/%d/%Y}"; ws["A2"].font = Font(bold=True, size=12, color="C00000")
    cols = ["Origem", "PO / Lote", "Box Mark", "Cliente", "Produto", "Caixa", "Qtd cx", "Bunches/cx", "Unid./bunch",
            "Total unid.", "Custo x U", "Total $", "Confirmado", "Notas"]
    r = 4
    for c, n in enumerate(cols, 1):
        cell = ws.cell(r, c, n); cell.font = Font(bold=True, color="FFFFFF"); cell.border = box
        cell.fill = PatternFill("solid", fgColor="1F3864"); cell.alignment = Alignment(horizontal="center")
    tq = tt = 0
    for l in sorted(lines, key=lambda l: (l["origem"] != "Purchase Control", l["cliente"], l["produto"])):
        r += 1; tot = round(l["total"] * l["preco"], 2); tq += l["qtd"]; tt += tot
        vals = [l["origem"], l["ref"], l["fp"], l["cliente"], l["produto"], l["caixa"], l["qtd"], l["bunches"], l["ux"],
                l["total"], l["preco"], tot, "SIM" if l["confirmado"] else "NÃO", l["notas"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell.border = box
            if c in (11, 12): cell.number_format = "#,##0.00"
            if 6 <= c <= 13: cell.alignment = Alignment(horizontal="center")
    r += 1
    ws.cell(r, 5, "TOTAL").font = Font(bold=True); ws.cell(r, 7, tq).font = Font(bold=True)
    ws.cell(r, 12, round(tt, 2)).font = Font(bold=True); ws.cell(r, 12).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJKLMN", (16, 26, 11, 30, 34, 10, 8, 11, 11, 10, 10, 10, 11, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"; wb.save(path)
    return tq, round(tt, 2)

def next_friday(today):
    return today + timedelta(days=(4 - today.weekday()) % 7 or 7)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date"); ap.add_argument("--dry-run", action="store_true"); ap.add_argument("--test", action="store_true")
    ap.add_argument("--out", default="out")
    a = ap.parse_args()
    now = datetime.now(MIAMI)
    day = date.fromisoformat(a.date) if a.date else next_friday(now.date())
    print(f"Hoje (Miami): {now:%a %Y-%m-%d %H:%M} → pedidos de {day:%a %Y-%m-%d}")
    orders = collect(day)
    os.makedirs(a.out, exist_ok=True)
    files, resumo, vazias = [], [], []
    for pref, (_, farm) in FARMS.items():
        lines = orders[pref]
        if not lines: vazias.append(farm); continue
        path = os.path.join(a.out, f"{farm} - {day:%b} {ordinal(day.day)}.xlsx")
        q, t = build_xlsx(path, farm, day, lines); files.append(path)
        npc = sum(1 for l in lines if l["origem"] == "Purchase Control")
        resumo.append(f"{farm}: {q} cx, ${t:,.2f} ({npc} do Purchase Control, {len(lines)-npc} do Inventory Entry)")
        print("  " + resumo[-1] + f" → {path}")
    if vazias: print("  sem pedido: " + ", ".join(vazias))
    if a.dry_run or not files: return
    subject = f"California farms for {day:%A} {day:%b} {ordinal(day.day)}"
    body = ("<p>Valerii,</p><p>Seguem os pedidos das fazendas para " + f"{day:%m/%d}" + ", um arquivo por fazenda:</p><ul>" +
            "".join(f"<li>{s}</li>" for s in resumo) + "</ul>" +
            (f"<p>Sem pedido: {', '.join(vazias)}.</p>" if vazias else ""))
    to, cc = (BUYER_CC, []) if a.test else (TO, BUYER_CC)
    if a.test: subject = "TESTE — " + subject + f" (iria para: {', '.join(TO)})"
    atts = []
    for p in files:
        with open(p, "rb") as f: atts.append({"filename": os.path.basename(p), "content": base64.b64encode(f.read()).decode()})
    r = requests.post("https://api.resend.com/emails", timeout=60, headers={"Authorization": f"Bearer {os.environ['RESEND_API_KEY']}"},
                      json={"from": FROM_EMAIL, "to": to, "cc": cc, "subject": subject, "html": body, "attachments": atts})
    if r.status_code >= 300: sys.exit(f"Resend {r.status_code}: {r.text[:200]}")
    print(f"  enviado ({r.json().get('id')}) para {to} cc {cc}")

if __name__ == "__main__":
    main()
